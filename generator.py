"""Generate formatted English Title Page as .docx for journal submission."""

from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CENTER = WD_ALIGN_PARAGRAPH.CENTER


def set_font(run, name="Times New Roman", size=11, bold=False):
    """Set font properties on a run."""
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


SUPERSCRIPT_SIZE = 13  # Large, visible superscript numbers


def add_superscript_run(paragraph, text, size=SUPERSCRIPT_SIZE):
    """Add superscript text (for affiliation numbers and asterisks)."""
    run = paragraph.add_run(text)
    run.font.superscript = True
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return run


def generate_docx(data):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.15

    # --- 1. Title (centered, bold) ---
    p_title = doc.add_paragraph()
    p_title.alignment = CENTER
    p_title.paragraph_format.space_after = Pt(16)
    run_label = p_title.add_run("Title：")
    set_font(run_label, bold=True, size=12)
    run_text = p_title.add_run(data.get("title", ""))
    set_font(run_text, size=12)

    # --- 2. Fund Information (left-aligned) ---
    fund_info = data.get("fund_info", [])
    if fund_info:
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(8)

        p_fund_label = doc.add_paragraph()
        p_fund_label.paragraph_format.space_after = Pt(4)
        run_fl = p_fund_label.add_run("Fund Information：")
        set_font(run_fl, bold=True, size=11)

        for line in fund_info:
            p_fund = doc.add_paragraph()
            p_fund.paragraph_format.space_after = Pt(2)
            run_f = p_fund.add_run(line)
            set_font(run_f, size=10.5)

        p_spacer2 = doc.add_paragraph()
        p_spacer2.paragraph_format.space_after = Pt(6)

    # --- 3. Authors line (centered) ---
    p_authors = doc.add_paragraph()
    p_authors.alignment = CENTER
    p_authors.paragraph_format.space_after = Pt(12)

    authors = data.get("authors", [])
    for i, author in enumerate(authors):
        if i > 0:
            run_sep = p_authors.add_run("    ")
            set_font(run_sep, size=11)
        run_name = p_authors.add_run(author["name"])
        set_font(run_name, size=11)
        aff_nums = author.get("affiliations", [])
        for j, num in enumerate(aff_nums):
            if j > 0:
                run_comma = p_authors.add_run(",")
                run_comma.font.superscript = True
                run_comma.font.size = Pt(SUPERSCRIPT_SIZE)
                run_comma.font.name = "Times New Roman"
            add_superscript_run(p_authors, str(num))
        if author.get("is_corresponding"):
            add_superscript_run(p_authors, "*")

    # --- 4. Affiliations (centered, each line) ---
    affiliations = data.get("affiliations", [])
    for aff in affiliations:
        aff_lines = aff.get("lines", [])
        for j, line in enumerate(aff_lines):
            p_aff = doc.add_paragraph()
            p_aff.alignment = CENTER
            p_aff.paragraph_format.space_after = Pt(0)
            p_aff.paragraph_format.space_before = Pt(0)
            if j == 0:
                run_num = p_aff.add_run(f"{aff['num']}")
                set_font(run_num, size=10.5)
                run_rest = p_aff.add_run(f" {line}")
                set_font(run_rest, size=10.5)
            else:
                run_rest = p_aff.add_run(line)
                set_font(run_rest, size=10.5)

        p_aff_spacer = doc.add_paragraph()
        p_aff_spacer.paragraph_format.space_after = Pt(4)

    # --- 5. Corresponding author declaration (left-aligned) ---
    corr_name = ""
    for a in authors:
        if a.get("is_corresponding"):
            corr_name = a["name"]
            break

    p_corr = doc.add_paragraph()
    p_corr.paragraph_format.space_after = Pt(12)
    run_c1 = p_corr.add_run('"*" means corresponding author: ')
    set_font(run_c1, size=10.5)
    run_c2 = p_corr.add_run(corr_name)
    set_font(run_c2, size=10.5)

    # --- 6. Email (left-aligned) ---
    p_email_label = doc.add_paragraph()
    p_email_label.paragraph_format.space_after = Pt(4)
    run_el = p_email_label.add_run("Email:")
    set_font(run_el, bold=True, size=11)

    for author in authors:
        email = author.get("email", "").strip()
        if email:
            p_email = doc.add_paragraph()
            p_email.paragraph_format.space_after = Pt(1)
            run_e = p_email.add_run(f"{email}  ({author['name']})")
            set_font(run_e, size=10.5)

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
