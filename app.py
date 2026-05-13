"""Flask app for SCI Title Page Generator."""

import os
import tempfile
import subprocess

from flask import Flask, render_template, request, jsonify, send_file
from generator import generate_docx

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024  # 32MB max upload

ALLOWED_EXTENSIONS = {".doc", ".docx"}


def extract_docx_text(filepath):
    """Extract text from .docx file."""
    from docx import Document
    doc = Document(filepath)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_doc_text(filepath):
    """Extract text from .doc file using macOS textutil."""
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        subprocess.run(
            ["textutil", "-convert", "txt", filepath, "-output", tmp_path],
            check=True, capture_output=True, timeout=30,
        )
        with open(tmp_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return text
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/journal-filter")
def journal_filter():
    return render_template("journal-filter.html")


@app.route("/api/generate", methods=["POST"])
def generate():
    data = request.get_json()
    if not data:
        return jsonify({"error": "No data provided"}), 400

    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title is required"}), 400

    filename = data.get("filename", "title_page").strip()
    if not filename:
        filename = "title_page"

    normalized = {
        "title": data.get("title", ""),
        "fund_info": data.get("fundInfo", []),
        "authors": [],
        "affiliations": [],
    }
    for a in data.get("authors", []):
        normalized["authors"].append({
            "name": a.get("name", ""),
            "affiliations": a.get("affiliations", []),
            "is_corresponding": a.get("isCorresponding", False),
            "email": a.get("email", ""),
            "title": a.get("title", ""),
        })
    for aff in data.get("affiliations", []):
        normalized["affiliations"].append({
            "num": aff.get("num", 0),
            "lines": aff.get("lines", []),
        })

    buf = generate_docx(normalized)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"{filename}.docx",
    )


@app.route("/api/verify-school", methods=["POST"])
def verify_school():
    """Return a search URL for verifying school/institution name."""
    data = request.get_json()
    query = data.get("query", "").strip()
    if not query:
        return jsonify({"error": "No query provided"}), 400
    search_url = f"https://www.bing.com/search?q={query}+institution+official+website"
    return jsonify({"search_url": search_url})


@app.route("/api/upload", methods=["POST"])
def upload():
    """Upload a .doc or .docx file and extract text."""
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename"}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        return jsonify({"error": f"Unsupported format: {ext}. Only .doc and .docx are accepted."}), 400

    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        if ext == ".docx":
            text = extract_docx_text(tmp_path)
        else:
            text = extract_doc_text(tmp_path)

        return jsonify({"text": text, "filename": file.filename})
    except Exception as e:
        return jsonify({"error": f"Failed to extract text: {str(e)}"}), 500
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


if __name__ == "__main__":
    app.run(debug=True, port=5050)
